from flask import Blueprint, Flask, request, jsonify, redirect, url_for, render_template,session, Response, send_file
import io
from datetime import datetime
from flask_login import login_user, login_required, logout_user
import os
import json
import traceback
import utils.config
from utils.jsonwfc_object import JSONWFCObject
import utils.string_util
from utils.exception_util import ValidationError, NotFoundError, DBError
import resources.messages


from app.service.entriesinitapi.entriesinitapi_service import EntriesinitapiService
from app.service.entrydetailinitapi.entrydetailinitapi_service import EntrydetailinitapiService
from app.service.entryforminitapi.entryforminitapi_service import EntryforminitapiService
from app.service.entryformnewinitapi.entryformnewinitapi_service import EntryformnewinitapiService
from app.service.entrysaveapi.entrysaveapi_service import EntrysaveapiService
from app.service.entryupdateapi.entryupdateapi_service import EntryupdateapiService
from app.service.dashboardinitapi.dashboardinitapi_service import DashboardinitapiService
from app.service.dashboardnationalinitapi.dashboardnationalinitapi_service import DashboardnationalinitapiService
from app.service.dashboardprefinitapi.dashboardprefinitapi_service import DashboardprefinitapiService
from app.service.dashboardshokokaiinitapi.dashboardshokokaiinitapi_service import DashboardshokokaiinitapiService
from app.service.loginapi.loginapi_service import LoginapiService
from app.service.logininitapi.logininitapi_service import LogininitapiService
from app.service.logoutapi.logoutapi_service import LogoutapiService
from app.service.verify2faapi.verify2faapi_service import Verify2faapiService
from app.dto.entriesinitapi.entriesinitapi_dto import EntriesinitapiDto
from app.dto.entrydetailinitapi.entrydetailinitapi_dto import EntrydetailinitapiDto
from app.dto.entryforminitapi.entryforminitapi_dto import EntryforminitapiDto
from app.dto.entryformnewinitapi.entryformnewinitapi_dto import EntryformnewinitapiDto
from app.dto.entrysaveapi.entrysaveapi_dto import EntrysaveapiDto
from app.dto.entryupdateapi.entryupdateapi_dto import EntryupdateapiDto
from app.dto.dashboardinitapi.dashboardinitapi_dto import DashboardinitapiDto
from app.dto.dashboardnationalinitapi.dashboardnationalinitapi_dto import DashboardnationalinitapiDto
from app.dto.dashboardprefinitapi.dashboardprefinitapi_dto import DashboardprefinitapiDto
from app.dto.dashboardshokokaiinitapi.dashboardshokokaiinitapi_dto import DashboardshokokaiinitapiDto
from app.dto.loginapi.loginapi_dto import LoginapiDto
from app.dto.logininitapi.logininitapi_dto import LogininitapiDto
from app.dto.logoutapi.logoutapi_dto import LogoutapiDto
from app.dto.verify2faapi.verify2faapi_dto import Verify2faapiDto
from app.service.vectordetailinitapi.vectordetailinitapi_service import VectordetailinitapiService
from app.service.vectorforminitapi.vectorforminitapi_service import VectorforminitapiService
from app.service.vectorformnewinitapi.vectorformnewinitapi_service import VectorformnewinitapiService
from app.service.vectorsaveapi.vectorsaveapi_service import VectorsaveapiService
from app.service.vectorsinitapi.vectorsinitapi_service import VectorsinitapiService
from app.service.vectorsresyncapi.vectorsresyncapi_service import VectorsresyncapiService
from app.service.vectorupdateapi.vectorupdateapi_service import VectorupdateapiService
from app.service.monthlyinitapi.monthlyinitapi_service import MonthlyinitapiService
from app.service.monthlyfiscalyearapi.monthlyfiscalyearapi_service import MonthlyfiscalyearapiService
from app.service.monthlyexportapi.monthlyexportapi_service import MonthlyexportapiService
from app.service.monthlycsvexportapi.monthlycsvexportapi_service import MonthlycsvexportapiService
from app.dto.vectordetailinitapi.vectordetailinitapi_dto import VectordetailinitapiDto
from app.dto.vectorforminitapi.vectorforminitapi_dto import VectorforminitapiDto
from app.dto.vectorformnewinitapi.vectorformnewinitapi_dto import VectorformnewinitapiDto
from app.dto.vectorsaveapi.vectorsaveapi_dto import VectorsaveapiDto
from app.dto.vectorsinitapi.vectorsinitapi_dto import VectorsinitapiDto
from app.dto.vectorsresyncapi.vectorsresyncapi_dto import VectorsresyncapiDto
from app.dto.vectorupdateapi.vectorupdateapi_dto import VectorupdateapiDto
from app.dto.monthlyinitapi.monthlyinitapi_dto import MonthlyinitapiDto
from app.dto.monthlyfiscalyearapi.monthlyfiscalyearapi_dto import MonthlyfiscalyearapiDto
from app.dto.monthlyexportapi.monthlyexportapi_dto import MonthlyexportapiDto
from app.dto.monthlycsvexportapi.monthlycsvexportapi_dto import MonthlycsvexportapiDto
from app.service.aiformatapi.aiformatapi_service import AiformatapiService
from app.service.draftsaveapi.draftsaveapi_service import DraftsaveapiService
from app.service.voiceextendapi.voiceextendapi_service import VoiceextendapiService
from app.service.voiceinsertcontentapi.voiceinsertcontentapi_service import VoiceinsertcontentapiService
from app.service.voicerecordapi.voicerecordapi_service import VoicerecordapiService
from app.service.voiceuploadapi.voiceuploadapi_service import VoiceuploadapiService
from app.dto.aiformatapi.aiformatapi_dto import AiformatapiDto
from app.dto.aiinputinitapi.aiinputinitapi_dto import AiinputinitapiDto
from app.dto.draftsaveapi.draftsaveapi_dto import DraftsaveapiDto
from app.dto.voiceextendapi.voiceextendapi_dto import VoiceextendapiDto
from app.dto.voiceinsertcontentapi.voiceinsertcontentapi_dto import VoiceinsertcontentapiDto
from app.dto.voicerecordapi.voicerecordapi_dto import VoicerecordapiDto
from app.dto.voiceuploadapi.voiceuploadapi_dto import VoiceuploadapiDto
from app.service.knowledgeinitapi.knowledgeinitapi_service import KnowledgeinitapiService
from app.service.documentdetailinitapi.documentdetailinitapi_service import DocumentdetailinitapiService
from app.service.documentforminitapi.documentforminitapi_service import DocumentforminitapiService
from app.service.documentformnewinitapi.documentformnewinitapi_service import DocumentformnewinitapiService
from app.service.documentsaveapi.documentsaveapi_service import DocumentsaveapiService
from app.service.documentupdateapi.documentupdateapi_service import DocumentupdateapiService
from app.service.documentversiondownloadapi.documentversiondownloadapi_service import DocumentversiondownloadapiService
from app.service.documentversionnewapi.documentversionnewapi_service import DocumentversionnewapiService
from app.dto.knowledgeinitapi.knowledgeinitapi_dto import KnowledgeinitapiDto
from app.dto.documentdetailinitapi.documentdetailinitapi_dto import DocumentdetailinitapiDto
from app.dto.documentforminitapi.documentforminitapi_dto import DocumentforminitapiDto
from app.dto.documentformnewinitapi.documentformnewinitapi_dto import DocumentformnewinitapiDto
from app.dto.documentsaveapi.documentsaveapi_dto import DocumentsaveapiDto
from app.dto.documentupdateapi.documentupdateapi_dto import DocumentupdateapiDto
from app.dto.documentversiondownloadapi.documentversiondownloadapi_dto import DocumentversiondownloadapiDto
from app.dto.documentversionnewapi.documentversionnewapi_dto import DocumentversionnewapiDto

from app.dto.accounteditinitapi.accounteditinitapi_dto import AccounteditinitapiDto
from app.dto.accountforminitapi.accountforminitapi_dto import AccountforminitapiDto
from app.dto.accountsaveapi.accountsaveapi_dto import AccountsaveapiDto
from app.dto.accountsdetailapi.accountsdetailapi_dto import AccountsdetailapiDto
from app.dto.accountsfilterapi.accountsfilterapi_dto import AccountsfilterapiDto
from app.dto.accountsinitapi.accountsinitapi_dto import AccountsinitapiDto
from app.dto.accountupdateapi.accountupdateapi_dto import AccountupdateapiDto
from app.dto.accountdeleteapi.accountdeleteapi_dto import AccountdeleteapiDto
from app.service.accounteditinitapi.accounteditinitapi_service import AccounteditinitapiService
from app.service.accountforminitapi.accountforminitapi_service import AccountforminitapiService
from app.service.accountsaveapi.accountsaveapi_service import AccountsaveapiService
from app.service.accountsdetailapi.accountsdetailapi_service import AccountsdetailapiService
from app.service.accountsfilterapi.accountsfilterapi_service import AccountsfilterapiService
from app.service.accountsinitapi.accountsinitapi_service import AccountsinitapiService
from app.service.accountupdateapi.accountupdateapi_service import AccountupdateapiService
from app.service.accountdeleteapi.accountdeleteapi_service import AccountdeleteapiService
from app.dto.aiproposalgenerateapi.aiproposalgenerateapi_dto import AiproposalgenerateapiDto
from app.service.aiproposalgenerateapi.aiproposalgenerateapi_service import AiproposalgenerateapiService
from app.dto.aiproposalinitapi.aiproposalinitapi_dto import AiproposalinitapiDto
from app.service.aiproposalinitapi.aiproposalinitapi_service import AiproposalinitapiService
from app.dto.aiproposalregenerateapi.aiproposalregenerateapi_dto import AiproposalregenerateapiDto
from app.service.aiproposalregenerateapi.aiproposalregenerateapi_service import AiproposalregenerateapiService
from app.dto.knowledgesearchapi.knowledgesearchapi_dto import KnowledgesearchapiDto
from app.service.knowledgesearchapi.knowledgesearchapi_service import KnowledgesearchapiService
from app.dto.dashboardheatmapcelltoggleapi.dashboardheatmapcelltoggleapi_dto import DashboardheatmapcelltoggleapiDto
from app.service.dashboardheatmapcelltoggleapi.dashboardheatmapcelltoggleapi_service import DashboardheatmapcelltoggleapiService
from app.dto.dashboardheatmapfilterapi.dashboardheatmapfilterapi_dto import DashboardheatmapfilterapiDto
from app.service.dashboardheatmapfilterapi.dashboardheatmapfilterapi_service import DashboardheatmapfilterapiService
from app.dto.dashboardheatmappageapi.dashboardheatmappageapi_dto import DashboardheatmappageapiDto
from app.service.dashboardheatmappageapi.dashboardheatmappageapi_service import DashboardheatmappageapiService
from app.dto.draftsaveapi.draftsaveapi_dto import DraftsaveapiDto
from app.service.draftsaveapi.draftsaveapi_service import DraftsaveapiService
from app.dto.jigyoshomeinokohokakonosodanrirekihistoryapi.jigyoshomeinokohokakonosodanrirekihistoryapi_dto import JigyoshomeinokohokakonosodanrirekihistoryapiDto
from app.service.jigyoshomeinokohokakonosodanrirekihistoryapi.jigyoshomeinokohokakonosodanrirekihistoryapi_service import JigyoshomeinokohokakonosodanrirekihistoryapiService
from app.dto.jigyoshomeinokohokakonosodanrirekinamesapi.jigyoshomeinokohokakonosodanrirekinamesapi_dto import JigyoshomeinokohokakonosodanrirekinamesapiDto
from app.service.jigyoshomeinokohokakonosodanrirekinamesapi.jigyoshomeinokohokakonosodanrirekinamesapi_service import JigyoshomeinokohokakonosodanrirekinamesapiService
from app.dto.verify2fainitapi.verify2fainitapi_dto import Verify2fainitapiDto
from app.service.verify2fainitapi.verify2fainitapi_service import Verify2fainitapiService
from app.dto.formexportapi.formexportapi_dto import FormexportapiDto
from app.service.formexportapi.formexportapi_service import FormexportapiService
from app.dto.aisummarizeapi.aisummarizeapi_dto import AisummarizeapiDto
from app.service.aisummarizeapi.aisummarizeapi_service import AisummarizeapiService
from app.dto.formnewsaveapi.formnewsaveapi_dto import FormnewsaveapiDto
from app.service.formnewsaveapi.formnewsaveapi_service import FormnewsaveapiService
from app.dto.reportsinitapi.reportsinitapi_dto import ReportsinitapiDto
from app.service.reportsinitapi.reportsinitapi_service import ReportsinitapiService
from app.dto.reportssearchapi.reportssearchapi_dto import ReportssearchapiDto
from app.service.reportssearchapi.reportssearchapi_service import ReportssearchapiService
from app.dto.reportscsvexportapi.reportscsvexportapi_dto import ReportscsvexportapiDto
from app.service.reportscsvexportapi.reportscsvexportapi_service import ReportscsvexportapiService
from app.dto.reportdetailinitapi.reportdetailinitapi_dto import ReportdetailinitapiDto
from app.service.reportdetailinitapi.reportdetailinitapi_service import ReportdetailinitapiService


commonfunction_route = Blueprint('commonfunction_route', __name__)

#
# entriesinitapi - 知識データ一覧画面初期表示 - サーバー関数
# @param entriesinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entriesinitapi.do", methods=['POST'])

def entriesinitapi() :
	"""entriesinitapi - 知識データ一覧画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entriesinitapi_dto = EntriesinitapiDto.dict_to_json(data)
	entriesinitapiVar_service = EntriesinitapiService()

	entriesinitapiVar_service.entriesinitapi(entriesinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# entrydetailinitapi - 知識データ詳細画面初期表示 - サーバー関数
# @param entrydetailinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entrydetailinitapi.do", methods=['POST'])

def entrydetailinitapi() :
	"""entrydetailinitapi - 知識データ詳細画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entrydetailinitapi_dto = EntrydetailinitapiDto.dict_to_json(data)
	entrydetailinitapiVar_service = EntrydetailinitapiService()

	entrydetailinitapiVar_service.entrydetailinitapi(entrydetailinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# entryforminitapi - 知識データ編集編集 - サーバー関数
# @param entryforminitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entryforminitapi.do", methods=['POST'])

def entryforminitapi() :
	"""entryforminitapi - 知識データ編集編集 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entryforminitapi_dto = EntryforminitapiDto.dict_to_json(data)
	entryforminitapiVar_service = EntryforminitapiService()

	entryforminitapiVar_service.entryforminitapi(entryforminitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# entryformnewinitapi - 知識データ新規＋ 知識データを登録 - サーバー関数
# @param entryformnewinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entryformnewinitapi.do", methods=['POST'])

def entryformnewinitapi() :
	"""entryformnewinitapi - 知識データ新規＋ 知識データを登録 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entryformnewinitapi_dto = EntryformnewinitapiDto.dict_to_json(data)
	entryformnewinitapiVar_service = EntryformnewinitapiService()

	entryformnewinitapiVar_service.entryformnewinitapi(entryformnewinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# entrysaveapi - 知識データ新規画面登録ボタン - サーバー関数
# @param entrysaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entrysaveapi.do", methods=['POST'])

def entrysaveapi() :
	"""entrysaveapi - 知識データ新規画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entrysaveapi_dto = EntrysaveapiDto.dict_to_json(data)
	entrysaveapiVar_service = EntrysaveapiService()

	entrysaveapiVar_service.entrysaveapi(entrysaveapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# entryupdateapi - 知識データ編集画面登録ボタン - サーバー関数
# @param entryupdateapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/entryupdateapi.do", methods=['POST'])

def entryupdateapi() :
	"""entryupdateapi - 知識データ編集画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	entryupdateapi_dto = EntryupdateapiDto.dict_to_json(data)
	entryupdateapiVar_service = EntryupdateapiService()

	entryupdateapiVar_service.entryupdateapi(entryupdateapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# dashboardinitapi - ダッシュボード（商工会）ダッシュボードを開く（初期データ取 - サーバー関数
# @param dashboardinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/dashboardinitapi.do", methods=['POST'])

def dashboardinitapi() :
	"""dashboardinitapi - ダッシュボード（商工会）ダッシュボードを開く（初期データ取 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	dashboardinitapi_dto = DashboardinitapiDto.dict_to_json(data)
	dashboardinitapiVar_service = DashboardinitapiService()

	dashboardinitapiVar_service.dashboardinitapi(dashboardinitapi_dto, jsonObj)

	return jsonObj.toJsonString()


#
# dashboardnationalinitapi - ダッシュボード（全国連）初期データ取得
#

@commonfunction_route.route("/dashboardnationalinitapi.do", methods=['POST'])
def dashboardnationalinitapi():
	"""dashboardnationalinitapi - ダッシュボード（全国連）初期データ取得"""
	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()
	dto = DashboardnationalinitapiDto.dict_to_json(data)
	DashboardnationalinitapiService().dashboardnationalinitapi(dto, jsonObj)
	return jsonObj.toJsonString()


#
# dashboardprefinitapi - ダッシュボード（県連）初期データ取得
#

@commonfunction_route.route("/dashboardprefinitapi.do", methods=['POST'])
def dashboardprefinitapi():
	"""dashboardprefinitapi - ダッシュボード（県連）初期データ取得"""
	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()
	dto = DashboardprefinitapiDto.dict_to_json(data)
	DashboardprefinitapiService().dashboardprefinitapi(dto, jsonObj)
	return jsonObj.toJsonString()


#
# dashboardshokokaiinitapi - ダッシュボード（商工会）初期データ取得
#

@commonfunction_route.route("/dashboardshokokaiinitapi.do", methods=['POST'])
def dashboardshokokaiinitapi():
	"""dashboardshokokaiinitapi - ダッシュボード（商工会）初期データ取得"""
	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()
	dto = DashboardshokokaiinitapiDto.dict_to_json(data)
	DashboardshokokaiinitapiService().dashboardshokokaiinitapi(dto, jsonObj)
	return jsonObj.toJsonString()



#
# loginapi - ログイン画面ログインボタン - サーバー関数
# @param loginapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/loginapi.do", methods=['POST'])

def loginapi() :
	"""loginapi - ログイン画面ログインボタン - サーバー関数"""

	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	loginapi_dto = LoginapiDto.dict_to_json(data)
	loginapiVar_service = LoginapiService()

	loginapiVar_service.loginapi(loginapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# logininitapi - ログイン画面 - サーバー関数
# @param logininitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/logininitapi.do", methods=['POST'])

def logininitapi() :
	"""logininitapi - ログイン画面 - サーバー関数"""

	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	logininitapi_dto = LogininitapiDto.dict_to_json(data)
	logininitapiVar_service = LogininitapiService()

	logininitapiVar_service.logininitapi(logininitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# logoutapi - ログアウト - サーバー関数
# @param logoutapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/logoutapi.do", methods=['POST'])

def logoutapi() :
	"""logoutapi - ログアウト - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	logoutapi_dto = LogoutapiDto.dict_to_json(data)
	logoutapiVar_service = LogoutapiService()

	logoutapiVar_service.logoutapi(logoutapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# verify2faapi - 二要素認証画面認証ボタン - サーバー関数
# @param verify2faapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/verify2faapi.do", methods=['POST'])

def verify2faapi() :
	"""verify2faapi - 二要素認証画面認証ボタン - サーバー関数"""

	data = request.get_json(silent=True) or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	verify2faapi_dto = Verify2faapiDto.dict_to_json(data)
	verify2faapiVar_service = Verify2faapiService()

	verify2faapiVar_service.verify2faapi(verify2faapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectordetailinitapi - ベクトル詳細画面初期表示 - サーバー関数
# @param vectordetailinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectordetailinitapi.do", methods=['POST'])

def vectordetailinitapi() :
	"""vectordetailinitapi - ベクトル詳細画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectordetailinitapi_dto = VectordetailinitapiDto.dict_to_json(data)
	vectordetailinitapiVar_service = VectordetailinitapiService()

	vectordetailinitapiVar_service.vectordetailinitapi(vectordetailinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorforminitapi - ベクトルコレクション編集編集 - サーバー関数
# @param vectorforminitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorforminitapi.do", methods=['POST'])

def vectorforminitapi() :
	"""vectorforminitapi - ベクトルコレクション編集編集 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorforminitapi_dto = VectorforminitapiDto.dict_to_json(data)
	vectorforminitapiVar_service = VectorforminitapiService()

	vectorforminitapiVar_service.vectorforminitapi(vectorforminitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorformnewinitapi - ベクトルコレクション新規＋ コレクションを登録 - サーバー関数
# @param vectorformnewinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorformnewinitapi.do", methods=['POST'])

def vectorformnewinitapi() :
	"""vectorformnewinitapi - ベクトルコレクション新規＋ コレクションを登録 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorformnewinitapi_dto = VectorformnewinitapiDto.dict_to_json(data)
	vectorformnewinitapiVar_service = VectorformnewinitapiService()

	vectorformnewinitapiVar_service.vectorformnewinitapi(vectorformnewinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorsaveapi - ベクトル新規画面登録ボタン - サーバー関数
# @param vectorsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorsaveapi.do", methods=['POST'])

def vectorsaveapi() :
	"""vectorsaveapi - ベクトル新規画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorsaveapi_dto = VectorsaveapiDto.dict_to_json(data)
	vectorsaveapiVar_service = VectorsaveapiService()

	vectorsaveapiVar_service.vectorsaveapi(vectorsaveapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorsinitapi - ベクトル一覧画面初期表示 - サーバー関数
# @param vectorsinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorsinitapi.do", methods=['POST'])

def vectorsinitapi() :
	"""vectorsinitapi - ベクトル一覧画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorsinitapi_dto = VectorsinitapiDto.dict_to_json(data)
	vectorsinitapiVar_service = VectorsinitapiService()

	vectorsinitapiVar_service.vectorsinitapi(vectorsinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorsresyncapi - ベクトル一覧画面再同期ボタン - サーバー関数
# @param vectorsresyncapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorsresyncapi.do", methods=['POST'])

def vectorsresyncapi() :
	"""vectorsresyncapi - ベクトル一覧画面再同期ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorsresyncapi_dto = VectorsresyncapiDto.dict_to_json(data)
	vectorsresyncapiVar_service = VectorsresyncapiService()

	vectorsresyncapiVar_service.vectorsresyncapi(vectorsresyncapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# vectorupdateapi - ベクトル編集画面登録ボタン - サーバー関数
# @param vectorupdateapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/vectorupdateapi.do", methods=['POST'])

def vectorupdateapi() :
	"""vectorupdateapi - ベクトル編集画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	vectorupdateapi_dto = VectorupdateapiDto.dict_to_json(data)
	vectorupdateapiVar_service = VectorupdateapiService()

	vectorupdateapiVar_service.vectorupdateapi(vectorupdateapi_dto, jsonObj)

	return jsonObj.toJsonString()

# accounteditinitapi - アカウント編集画面初期表示 - サーバー関数
# @param accounteditinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accounteditinitapi.do", methods=['POST'])

def accounteditinitapi() :
	"""accounteditinitapi - アカウント編集画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accounteditinitapi_dto = AccounteditinitapiDto.dict_to_json(data)
	accounteditinitapiVar_service = AccounteditinitapiService()

	accounteditinitapiVar_service.accounteditinitapi(accounteditinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountforminitapi - アカウント新規画面初期表示 - サーバー関数
# @param accountforminitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountforminitapi.do", methods=['POST'])

def accountforminitapi() :
	"""accountforminitapi - アカウント新規画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountforminitapi_dto = AccountforminitapiDto.dict_to_json(data)
	accountforminitapiVar_service = AccountforminitapiService()

	accountforminitapiVar_service.accountforminitapi(accountforminitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountsaveapi - アカウント新規画面登録ボタン - サーバー関数
# @param accountsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountsaveapi.do", methods=['POST'])

def accountsaveapi() :
	"""accountsaveapi - アカウント新規画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountsaveapi_dto = AccountsaveapiDto.dict_to_json(data)
	accountsaveapiVar_service = AccountsaveapiService()

	accountsaveapiVar_service.accountsaveapi(accountsaveapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountsdetailapi - アカウント一覧画面詳細ボタン - サーバー関数
# @param accountsdetailapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountsdetailapi.do", methods=['POST'])

def accountsdetailapi() :
	"""accountsdetailapi - アカウント一覧画面詳細ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountsdetailapi_dto = AccountsdetailapiDto.dict_to_json(data)
	accountsdetailapiVar_service = AccountsdetailapiService()

	accountsdetailapiVar_service.accountsdetailapi(accountsdetailapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountsfilterapi - アカウント一覧一覧絞込（県/商工会/権限/検索） - サーバー関数
# @param accountsfilterapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountsfilterapi.do", methods=['POST'])

def accountsfilterapi() :
	"""accountsfilterapi - アカウント一覧一覧絞込（県/商工会/権限/検索） - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountsfilterapi_dto = AccountsfilterapiDto.dict_to_json(data)
	accountsfilterapiVar_service = AccountsfilterapiService()

	accountsfilterapiVar_service.accountsfilterapi(accountsfilterapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountsinitapi - アカウント一覧画面初期表示 - サーバー関数
# @param accountsinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountsinitapi.do", methods=['POST'])

def accountsinitapi() :
	"""accountsinitapi - アカウント一覧画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountsinitapi_dto = AccountsinitapiDto.dict_to_json(data)
	accountsinitapiVar_service = AccountsinitapiService()

	accountsinitapiVar_service.accountsinitapi(accountsinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountupdateapi - アカウント編集画面登録ボタン - サーバー関数
# @param accountupdateapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountupdateapi.do", methods=['POST'])

def accountupdateapi() :
	"""accountupdateapi - アカウント編集画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountupdateapi_dto = AccountupdateapiDto.dict_to_json(data)
	accountupdateapiVar_service = AccountupdateapiService()

	accountupdateapiVar_service.accountupdateapi(accountupdateapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# accountdeleteapi - アカウント詳細画面削除ボタン - サーバー関数
# @param accountdeleteapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/accountdeleteapi.do", methods=['POST'])

def accountdeleteapi() :
	"""accountdeleteapi - アカウント詳細画面削除ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	accountdeleteapi_dto = AccountdeleteapiDto.dict_to_json(data)
	accountdeleteapiVar_service = AccountdeleteapiService()

	accountdeleteapiVar_service.accountdeleteapi(accountdeleteapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# monthlyexportapi - 月次報告月次帳票出力 - サーバー関数
# @param monthlyexportapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/monthlyexportapi.do", methods=['POST'])

def monthlyexportapi() :
	"""monthlyexportapi - 月次報告月次帳票出力 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	monthlyexportapi_dto = MonthlyexportapiDto.dict_to_json(data)
	monthlyexportapiVar_service = MonthlyexportapiService()

	monthlyexportapiVar_service.monthlyexportapi(monthlyexportapi_dto, jsonObj)

	# 成功時は設定ディレクトリ上のファイルをそのまま返す（顧客アップロード想定）
	payload = jsonObj.getJsonObj()
	filepath = payload.pop("filepath", None) if isinstance(payload, dict) else None
	filename = payload.get("filename") if isinstance(payload, dict) else None
	if filepath and payload.get("e") in (None, "") and os.path.isfile(filepath):
		return send_file(
			filepath,
			as_attachment=True,
			download_name=filename or os.path.basename(filepath),
			mimetype="application/octet-stream",
		)

	# エラー時は JSON（Content-Type を明示し、FE が blob 誤ダウンロードしないようにする）
	return Response(
		jsonObj.toJsonString(),
		mimetype="application/json; charset=utf-8",
	)



#
# monthlycsvexportapi - 実績確認一覧CSV出力 - サーバー関数
# @param monthlycsvexportapi_dto（fiscalyearcode / 組織；条件に関係なく一覧）
# @param result
# @throws Exception
#

@commonfunction_route.route("/monthlycsvexportapi.do", methods=['POST'])

def monthlycsvexportapi() :
	"""monthlycsvexportapi - 実績確認一覧CSV出力（条件に関係なく・DB）"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	monthlycsvexportapi_dto = MonthlycsvexportapiDto.dict_to_json(data)
	MonthlycsvexportapiService().monthlycsvexportapi(monthlycsvexportapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# monthlyfiscalyearapi - 月次報告年度切替 - サーバー関数
# @param monthlyfiscalyearapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/monthlyfiscalyearapi.do", methods=['POST'])

def monthlyfiscalyearapi() :
	"""monthlyfiscalyearapi - 月次報告年度切替 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	monthlyfiscalyearapi_dto = MonthlyfiscalyearapiDto.dict_to_json(data)
	monthlyfiscalyearapiVar_service = MonthlyfiscalyearapiService()

	monthlyfiscalyearapiVar_service.monthlyfiscalyearapi(monthlyfiscalyearapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# monthlyinitapi - 月次報告画面初期表示 - サーバー関数
# @param monthlyinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/monthlyinitapi.do", methods=['POST'])

def monthlyinitapi() :
	"""monthlyinitapi - 月次報告画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	monthlyinitapi_dto = MonthlyinitapiDto.dict_to_json(data)
	monthlyinitapiVar_service = MonthlyinitapiService()

	monthlyinitapiVar_service.monthlyinitapi(monthlyinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# aiformatapi - 傾聴内容変換AI文字起こし AI整形 - サーバー関数
# @param aiformatapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/aiformatapi.do", methods=['POST'])

def aiformatapi() :
	"""aiformatapi - 傾聴内容変換AI文字起こし AI整形 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	aiformatapi_dto = AiformatapiDto.dict_to_json(data)
	aiformatapiVar_service = AiformatapiService()

	aiformatapiVar_service.aiformatapi(aiformatapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# aiinputinitapi - 傾聴内容変換AI画面初期表示 - サーバー関数
# @param aiinputinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/aiinputinitapi.do", methods=['POST'])

def aiinputinitapi() :
	"""aiinputinitapi - 傾聴内容変換AI画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	aiinputinitapi_dto = AiinputinitapiDto.dict_to_json(data)
	aiinputinitapiVar_service = AiinputinitapiService()

	aiinputinitapiVar_service.aiinputinitapi(aiinputinitapi_dto, jsonObj)

	return jsonObj.toJsonString()




#
# formexportapi - 傾聴内容変換AI帳票出力 - サーバー関数
# @param formexportapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/formexportapi.do", methods=['POST'])

def formexportapi() :
	"""formexportapi - 傾聴内容変換AI帳票出力 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	formexportapi_dto = FormexportapiDto.dict_to_json(data)
	formexportapiVar_service = FormexportapiService()

	formexportapiVar_service.formexportapi(formexportapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# voiceextendapi - 傾聴内容変換AI録音時間 ＋30分延長 - サーバー関数
# @param voiceextendapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/voiceextendapi.do", methods=['POST'])

def voiceextendapi() :
	"""voiceextendapi - 傾聴内容変換AI録音時間 ＋30分延長 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	voiceextendapi_dto = VoiceextendapiDto.dict_to_json(data)
	voiceextendapiVar_service = VoiceextendapiService()

	voiceextendapiVar_service.voiceextendapi(voiceextendapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# voiceinsertcontentapi - 傾聴内容変換AI文字起こしを内容欄に反映 - サーバー関数
# @param voiceinsertcontentapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/voiceinsertcontentapi.do", methods=['POST'])

def voiceinsertcontentapi() :
	"""voiceinsertcontentapi - 傾聴内容変換AI文字起こしを内容欄に反映 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	voiceinsertcontentapi_dto = VoiceinsertcontentapiDto.dict_to_json(data)
	voiceinsertcontentapiVar_service = VoiceinsertcontentapiService()

	voiceinsertcontentapiVar_service.voiceinsertcontentapi(voiceinsertcontentapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# voicerecordapi - 傾聴内容変換AI音声録音 開始／停止 - サーバー関数
# @param voicerecordapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/voicerecordapi.do", methods=['POST'])

def voicerecordapi() :
	"""voicerecordapi - 傾聴内容変換AI音声録音 開始／停止 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	voicerecordapi_dto = VoicerecordapiDto.dict_to_json(data)
	voicerecordapiVar_service = VoicerecordapiService()

	voicerecordapiVar_service.voicerecordapi(voicerecordapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# voiceuploadapi - 傾聴内容変換AI音声ファイル アップロード - サーバー関数
# @param voiceuploadapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/voiceuploadapi.do", methods=['POST'])

def voiceuploadapi() :
	"""voiceuploadapi - 傾聴内容変換AI音声ファイル アップロード - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	voiceuploadapi_dto = VoiceuploadapiDto.dict_to_json(data)
	voiceuploadapiVar_service = VoiceuploadapiService()

	voiceuploadapiVar_service.voiceuploadapi(voiceuploadapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentdetailinitapi - ナレッジ文書詳細画面初期表示 - サーバー関数
# @param documentdetailinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentdetailinitapi.do", methods=['POST'])

def documentdetailinitapi() :
	"""documentdetailinitapi - ナレッジ文書詳細画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentdetailinitapi_dto = DocumentdetailinitapiDto.dict_to_json(data)
	documentdetailinitapiVar_service = DocumentdetailinitapiService()

	documentdetailinitapiVar_service.documentdetailinitapi(documentdetailinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentforminitapi - ナレッジ文書編集編集 - サーバー関数
# @param documentforminitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentforminitapi.do", methods=['POST'])

def documentforminitapi() :
	"""documentforminitapi - ナレッジ文書編集編集 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentforminitapi_dto = DocumentforminitapiDto.dict_to_json(data)
	documentforminitapiVar_service = DocumentforminitapiService()

	documentforminitapiVar_service.documentforminitapi(documentforminitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentformnewinitapi - ナレッジ文書新規＋ 文書を登録 - サーバー関数
# @param documentformnewinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentformnewinitapi.do", methods=['POST'])

def documentformnewinitapi() :
	"""documentformnewinitapi - ナレッジ文書新規＋ 文書を登録 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentformnewinitapi_dto = DocumentformnewinitapiDto.dict_to_json(data)
	documentformnewinitapiVar_service = DocumentformnewinitapiService()

	documentformnewinitapiVar_service.documentformnewinitapi(documentformnewinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentsaveapi - ナレッジ文書新規画面登録ボタン - サーバー関数
# @param documentsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentsaveapi.do", methods=['POST'])

def documentsaveapi() :
	"""documentsaveapi - ナレッジ文書新規画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentsaveapi_dto = DocumentsaveapiDto.dict_to_json(data)
	documentsaveapiVar_service = DocumentsaveapiService()

	documentsaveapiVar_service.documentsaveapi(documentsaveapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentupdateapi - ナレッジ文書編集画面登録ボタン - サーバー関数
# @param documentupdateapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentupdateapi.do", methods=['POST'])

def documentupdateapi() :
	"""documentupdateapi - ナレッジ文書編集画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentupdateapi_dto = DocumentupdateapiDto.dict_to_json(data)
	documentupdateapiVar_service = DocumentupdateapiService()

	documentupdateapiVar_service.documentupdateapi(documentupdateapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# documentversiondownloadapi - ナレッジ文書詳細⬇ ダウンロード - サーバー関数
# @param documentversiondownloadapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentversiondownloadapi.do", methods=['POST'])

def documentversiondownloadapi() :
	"""documentversiondownloadapi - ナレッジ文書詳細⬇ ダウンロード - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentversiondownloadapi_dto = DocumentversiondownloadapiDto.dict_to_json(data)
	documentversiondownloadapiVar_service = DocumentversiondownloadapiService()

	documentversiondownloadapiVar_service.documentversiondownloadapi(documentversiondownloadapi_dto, jsonObj)

	# jsonObjからダウンロード情報を取得する
	import json as _json
	download_info = jsonObj.getJsonObj().get("downloadInfo", "{}")
	info = _json.loads(download_info) if download_info else {}

	file_path = info.get("file_path", "")
	title = info.get("title", "文書")
	version_number = info.get("version_number", 1)

	# ファイルパスから拡張子を取得する
	ext = file_path.split(".")[-1].lower() if "." in file_path else "txt"

	# ログインユーザーIDを取得する（FlaskセッションのUSER_IDキーから取得）
	login_id = session.get("USER_ID", "unknown") if session else "unknown"

	# システム日付（yyyyMMddHHmmssSSS）を生成する
	now = datetime.now()
	timestamp = now.strftime("%Y%m%d%H%M%S") + str(now.microsecond // 1000).zfill(3)

	# ダウンロードファイル名を生成する（タイトル+システム日付+ログインID+.拡張子）
	filename = f"{title}{timestamp}{login_id}.{ext}"

	# ファイル形式に応じてコンテンツを生成する
	buf = io.BytesIO()
	if ext == "pdf":
		# fpdf2で有効なPDFファイルを生成する
		from fpdf import FPDF
		pdf = FPDF()
		pdf.add_page()
		# 日本語フォントを使用するためTTFフォントを追加する
		font_path = None
		for candidate in [r"C:\Windows\Fonts\msgothic.ttc", r"C:\Windows\Fonts\msmincho.ttc", r"C:\Windows\Fonts\yuGothic.ttc"]:
			if os.path.exists(candidate):
				font_path = candidate
				break
		if font_path:
			pdf.add_font("JP", "", font_path, uni=True)
			pdf.set_font("JP", size=14)
		else:
			pdf.set_font("Helvetica", size=14)
		pdf.cell(0, 10, title, ln=True)
		pdf.ln(10)
		pdf.set_font("Helvetica", size=10)
		pdf.cell(0, 10, "Version: " + str(version_number), ln=True)
		pdf.cell(0, 10, "File: " + file_path, ln=True)
		pdf.output(buf)
	elif ext == "xlsx":
		# openpyxlでExcelファイルを生成する
		from openpyxl import Workbook
		wb = Workbook()
		ws = wb.active
		ws.title = "文書情報"
		ws["A1"] = "タイトル"
		ws["B1"] = title
		ws["A2"] = "版番号"
		ws["B2"] = version_number
		ws["A3"] = "ファイルパス"
		ws["B3"] = file_path
		wb.save(buf)
	elif ext == "docx" or ext == "doc":
		# python-docxでWordファイルを生成する
		from docx import Document
		doc = Document()
		doc.add_heading(title, 0)
		doc.add_paragraph("版番号: " + str(version_number))
		doc.add_paragraph("ファイルパス: " + file_path)
		doc.save(buf)
	else:
		# テキストファイルとして出力する
		content = f"タイトル: {title}\n版番号: {version_number}\nファイルパス: {file_path}\n"
		buf.write(content.encode("utf-8"))

	buf.seek(0)
	# ファイル名をカスタムヘッダーで別途渡す（日本語対応）
	import urllib.parse
	encoded_filename = urllib.parse.quote(filename, safe='')
	resp = send_file(buf, as_attachment=True, download_name=filename, mimetype='application/octet-stream')
	resp.headers['X-Download-Filename'] = encoded_filename
	return resp



#
# documentversionnewapi - ナレッジ文書詳細画面版登録 - サーバー関数
# @param documentversionnewapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/documentversionnewapi.do", methods=['POST'])

def documentversionnewapi() :
	"""documentversionnewapi - ナレッジ文書詳細画面版登録 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	documentversionnewapi_dto = DocumentversionnewapiDto.dict_to_json(data)
	documentversionnewapiVar_service = DocumentversionnewapiService()

	documentversionnewapiVar_service.documentversionnewapi(documentversionnewapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# knowledgeinitapi - ナレッジ文書一覧画面初期表示 - サーバー関数
# @param knowledgeinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/knowledgeinitapi.do", methods=['POST'])

def knowledgeinitapi() :
	"""knowledgeinitapi - ナレッジ文書一覧画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	knowledgeinitapi_dto = KnowledgeinitapiDto.dict_to_json(data)
	knowledgeinitapiVar_service = KnowledgeinitapiService()

	knowledgeinitapiVar_service.knowledgeinitapi(knowledgeinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# aisummarizeapi - 報告書編集内容 AI要約→概要反映 - サーバー関数
# @param aisummarizeapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/aisummarizeapi.do", methods=['POST'])

def aisummarizeapi() :
	"""aisummarizeapi - 報告書編集内容 AI要約→概要反映 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	aisummarizeapi_dto = AisummarizeapiDto.dict_to_json(data)
	aisummarizeapiVar_service = AisummarizeapiService()

	aisummarizeapiVar_service.aisummarizeapi(aisummarizeapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# formeditinitapi - 報告書編集画面初期表示 - サーバー関数
# @param formeditinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/formeditinitapi.do", methods=['POST'])

def formeditinitapi() :
	"""formeditinitapi - 報告書編集画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	formeditinitapi_dto = FormeditinitapiDto.dict_to_json(data)
	formeditinitapiVar_service = FormeditinitapiService()

	formeditinitapiVar_service.formeditinitapi(formeditinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# formeditsaveapi - 報告書編集画面登録ボタン - サーバー関数
# @param formeditsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/formeditsaveapi.do", methods=['POST'])

def formeditsaveapi() :
	"""formeditsaveapi - 報告書編集画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	formeditsaveapi_dto = FormeditsaveapiDto.dict_to_json(data)
	formeditsaveapiVar_service = FormeditsaveapiService()

	formeditsaveapiVar_service.formeditsaveapi(formeditsaveapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# formnewinitapi - 報告書新規画面初期表示 - サーバー関数
# @param formnewinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/formnewinitapi.do", methods=['POST'])

def formnewinitapi() :
	"""formnewinitapi - 報告書新規画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	formnewinitapi_dto = FormnewinitapiDto.dict_to_json(data)
	formnewinitapiVar_service = FormnewinitapiService()

	formnewinitapiVar_service.formnewinitapi(formnewinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# formnewsaveapi - 報告書新規画面登録ボタン - サーバー関数
# @param formnewsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/formnewsaveapi.do", methods=['POST'])

def formnewsaveapi() :
	"""formnewsaveapi - 報告書新規画面登録ボタン - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	formnewsaveapi_dto = FormnewsaveapiDto.dict_to_json(data)
	formnewsaveapiVar_service = FormnewsaveapiService()

	formnewsaveapiVar_service.formnewsaveapi(formnewsaveapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# reportdetailinitapi - 報告書詳細画面初期表示 - サーバー関数
# @param reportdetailinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/reportdetailinitapi.do", methods=['POST'])

def reportdetailinitapi() :
	"""reportdetailinitapi - 報告書詳細画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	reportdetailinitapi_dto = ReportdetailinitapiDto.dict_to_json(data)
	reportdetailinitapiVar_service = ReportdetailinitapiService()

	reportdetailinitapiVar_service.reportdetailinitapi(reportdetailinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# reportscsvexportapi - 報告書一覧（様式F）CSV出力 - サーバー関数
# @param reportscsvexportapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/reportscsvexportapi.do", methods=['POST'])

def reportscsvexportapi() :
	"""reportscsvexportapi - 報告書一覧（様式F）CSV出力 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	reportscsvexportapi_dto = ReportscsvexportapiDto.dict_to_json(data)
	reportscsvexportapiVar_service = ReportscsvexportapiService()

	reportscsvexportapiVar_service.reportscsvexportapi(reportscsvexportapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# reportsinitapi - 報告書一覧画面初期表示 - サーバー関数
# @param reportsinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/reportsinitapi.do", methods=['POST'])

def reportsinitapi() :
	"""reportsinitapi - 報告書一覧画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	reportsinitapi_dto = ReportsinitapiDto.dict_to_json(data)
	reportsinitapiVar_service = ReportsinitapiService()

	reportsinitapiVar_service.reportsinitapi(reportsinitapi_dto, jsonObj)

	return jsonObj.toJsonString()



#
# reportssearchapi - 報告書一覧画面検索 - サーバー関数
# @param reportssearchapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/reportssearchapi.do", methods=['POST'])

def reportssearchapi() :
	"""reportssearchapi - 報告書一覧画面検索 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371


#UnitedControllerBuilder 963

	reportssearchapi_dto = ReportssearchapiDto.dict_to_json(data)
	reportssearchapiVar_service = ReportssearchapiService()

	reportssearchapiVar_service.reportssearchapi(reportssearchapi_dto, jsonObj)

	return jsonObj.toJsonString()

@commonfunction_route.route("/aiproposalgenerateapi.do", methods=['POST'])

def aiproposalgenerateapi() :
	"""aiproposalgenerateapi - AI支援提案AI提案を生成 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	aiproposalgenerateapi_dto = AiproposalgenerateapiDto.dict_to_json(data)
	aiproposalgenerateapiVar_service = AiproposalgenerateapiService()

	aiproposalgenerateapiVar_service.aiproposalgenerateapi(aiproposalgenerateapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# aiproposalinitapi - AI支援提案画面初期表示 - サーバー関数
# @param aiproposalinitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/aiproposalinitapi.do", methods=['POST'])

def aiproposalinitapi() :
	"""aiproposalinitapi - AI支援提案画面初期表示 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	aiproposalinitapi_dto = AiproposalinitapiDto.dict_to_json(data)
	aiproposalinitapiVar_service = AiproposalinitapiService()

	aiproposalinitapiVar_service.aiproposalinitapi(aiproposalinitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# aiproposalregenerateapi - AI支援提案違う内容を見る - サーバー関数
# @param aiproposalregenerateapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/aiproposalregenerateapi.do", methods=['POST'])

def aiproposalregenerateapi() :
	"""aiproposalregenerateapi - AI支援提案違う内容を見る - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	aiproposalregenerateapi_dto = AiproposalregenerateapiDto.dict_to_json(data)
	aiproposalregenerateapiVar_service = AiproposalregenerateapiService()

	aiproposalregenerateapiVar_service.aiproposalregenerateapi(aiproposalregenerateapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# knowledgesearchapi - AI支援提案ナレッジを検索 - サーバー関数
# @param knowledgesearchapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/knowledgesearchapi.do", methods=['POST'])

def knowledgesearchapi() :
	"""knowledgesearchapi - AI支援提案ナレッジを検索 - サーバー関数"""

	data = request.get_json()
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	knowledgesearchapi_dto = KnowledgesearchapiDto.dict_to_json(data)
	knowledgesearchapiVar_service = KnowledgesearchapiService()

	knowledgesearchapiVar_service.knowledgesearchapi(knowledgesearchapi_dto, jsonObj)
	
	return jsonObj.toJsonString()
@commonfunction_route.route("/dashboardheatmapcelltoggleapi.do", methods=['POST'])

def dashboardheatmapcelltoggleapi() :
	"""dashboardheatmapcelltoggleapi - ダッシュボード（全国連）ヒートマップセル除外トグル - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	dashboardheatmapcelltoggleapi_dto = DashboardheatmapcelltoggleapiDto.dict_to_json(data)
	dashboardheatmapcelltoggleapiVar_service = DashboardheatmapcelltoggleapiService()

	dashboardheatmapcelltoggleapiVar_service.dashboardheatmapcelltoggleapi(dashboardheatmapcelltoggleapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# dashboardheatmapfilterapi - ダッシュボード（全国連）ヒートマップ絞込（地方／グループ） - サーバー関数
# @param dashboardheatmapfilterapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/dashboardheatmapfilterapi.do", methods=['POST'])

def dashboardheatmapfilterapi() :
	"""dashboardheatmapfilterapi - ダッシュボード（全国連）ヒートマップ絞込（地方／グループ） - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	dashboardheatmapfilterapi_dto = DashboardheatmapfilterapiDto.dict_to_json(data)
	dashboardheatmapfilterapiVar_service = DashboardheatmapfilterapiService()

	dashboardheatmapfilterapiVar_service.dashboardheatmapfilterapi(dashboardheatmapfilterapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# dashboardheatmappageapi - ダッシュボード（全国連）ヒートマップページ送り - サーバー関数
# @param dashboardheatmappageapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/dashboardheatmappageapi.do", methods=['POST'])

def dashboardheatmappageapi() :
	"""dashboardheatmappageapi - ダッシュボード（全国連）ヒートマップページ送り - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	dashboardheatmappageapi_dto = DashboardheatmappageapiDto.dict_to_json(data)
	dashboardheatmappageapiVar_service = DashboardheatmappageapiService()

	dashboardheatmappageapiVar_service.dashboardheatmappageapi(dashboardheatmappageapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# draftsaveapi - 報告書編集下書き保存 / 傾聴内容変換AI下書き保存 - サーバー関数
# 両画面が同一エンドポイント・同一サービスを共用する
# @param draftsaveapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/draftsaveapi.do", methods=['POST'])

def draftsaveapi() :
	"""draftsaveapi - 報告書編集下書き保存 - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	draftsaveapi_dto = DraftsaveapiDto.dict_to_json(data)
	draftsaveapiVar_service = DraftsaveapiService()

	draftsaveapiVar_service.draftsaveapi(draftsaveapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# jigyoshomeinokohokakonosodanrirekihistoryapi - 事業所過去相談履歴 - サーバー関数
# @param jigyoshomeinokohokakonosodanrirekihistoryapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/jigyoshomeinokohokakonosodanrirekihistoryapi.do", methods=['POST'])

def jigyoshomeinokohokakonosodanrirekihistoryapi() :
	"""jigyoshomeinokohokakonosodanrirekihistoryapi - 事業所過去相談履歴 - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	jigyoshomeinokohokakonosodanrirekihistoryapi_dto = JigyoshomeinokohokakonosodanrirekihistoryapiDto.dict_to_json(data)
	jigyoshomeinokohokakonosodanrirekihistoryapiVar_service = JigyoshomeinokohokakonosodanrirekihistoryapiService()

	jigyoshomeinokohokakonosodanrirekihistoryapiVar_service.jigyoshomeinokohokakonosodanrirekihistoryapi(jigyoshomeinokohokakonosodanrirekihistoryapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# jigyoshomeinokohokakonosodanrirekinamesapi - 事業所名サジェスト - サーバー関数
# @param jigyoshomeinokohokakonosodanrirekinamesapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/jigyoshomeinokohokakonosodanrirekinamesapi.do", methods=['POST'])

def jigyoshomeinokohokakonosodanrirekinamesapi() :
	"""jigyoshomeinokohokakonosodanrirekinamesapi - 事業所名サジェスト - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	jigyoshomeinokohokakonosodanrirekinamesapi_dto = JigyoshomeinokohokakonosodanrirekinamesapiDto.dict_to_json(data)
	jigyoshomeinokohokakonosodanrirekinamesapiVar_service = JigyoshomeinokohokakonosodanrirekinamesapiService()

	jigyoshomeinokohokakonosodanrirekinamesapiVar_service.jigyoshomeinokohokakonosodanrirekinamesapi(jigyoshomeinokohokakonosodanrirekinamesapi_dto, jsonObj)
	
	return jsonObj.toJsonString()



#
# verify2fainitapi - 二段階認証画面画面を開く - サーバー関数
# @param verify2fainitapi_dto
# @param result
# @throws Exception
#

@commonfunction_route.route("/verify2fainitapi.do", methods=['POST'])

def verify2fainitapi() :
	"""verify2fainitapi - 二段階認証画面画面を開く - サーバー関数"""

	data = request.get_json() or {}
	jsonObj = JSONWFCObject()

	# validate parameter 371

	
#UnitedControllerBuilder 963

	verify2fainitapi_dto = Verify2fainitapiDto.dict_to_json(data)
	verify2fainitapiVar_service = Verify2fainitapiService()

	verify2fainitapiVar_service.verify2fainitapi(verify2fainitapi_dto, jsonObj)
	
	return jsonObj.toJsonString()




