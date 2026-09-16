#BasicService.vm
#make Service templete
# 傾聴内容変換AI帳票出力サービス
# 指定された報告書（report_id）または最新の報告書をDB（v_output_f_excel）から取得し、Excel出力する
import json
import os
import io
import tempfile
import utils.config
import threading
import utils.json_constant
from flask import session
from utils.jsonwfc_object import JSONWFCObject
from app.dao.api204_formexport.api204_formexport_dao import Api204FormexportDao
from app.dto.api204_formexport.api204_formexport_dto import Api204FormexportDto
from app.dto.formexportapi.formexportapi_dto import FormexportapiDto
import utils.string_util
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


class FormexportapiService :

	#	#
	# 傾聴内容変換AI帳票出力
	# @param Entity
	# @param jsonObj
	# @return Excelファイルパス（失敗時はNone）
	# @throws Exception
	#
	def formexportapi(self,formexportapi_dto,jsonObj) :

		#GeniusClientScript 1315
		REPORT_ID = formexportapi_dto.reportid  # 任意 GeninusClientScript 1318
		#UltimateGeniuBean 115
		utils.config.global_log.debug(str(threading.current_thread().native_id)+ ": start")
		try :
			#報告書新規帳票出力_FORMEXPORTAPI_(API)

			#「項目処理」（共通関数:FORMEXPORTAPI）,パラメータは（report_id任意）

			#以下の処理を行う。

			#関数「API204_FORMEXPORT」：帳票出力対象の報告書データを取得する（実データ：v_output_f_excel）。
			api204_formexport = Api204FormexportDto.dict_to_json({})
			api204_formexport.reportid = utils.string_util.changeNullToBlank(REPORT_ID)
			api204_formexportList = Api204FormexportDao().api204_formexport(api204_formexport)

			#取得結果（1件）を確認する。
			report_row = {}
			if api204_formexportList :
				report_row = api204_formexportList[0]

			if not report_row :
				jsonObj.setValue(utils.json_constant.JSONID_ERR, "出力対象の報告書が見つかりません。先に下書き保存または登録を行ってください")
				jsonObj.setValue(utils.json_constant.JSONID_FOR_RUNRESULT, utils.json_constant.RUNRESULT_FAIL)
				utils.config.global_log.debug(str(threading.current_thread().native_id)+ ": end")
				return None

			#Excel生成：v_output_f_excelの取得データから「F相談受付票」Excelファイルを出力する。
			file_path = self._generate_excel(report_row)

			#出力ファイル名を設定する。
			report_code = utils.string_util.changeNullToBlank(report_row.get("報告書番号", ""))
			file_name = "F_相談受付票_" + report_code + ".xlsx"
			jsonObj.setValue("dragFileName", file_name)
			jsonObj.setValue(utils.json_constant.JSONID_MSG, "帳票出力が完了しました")
			jsonObj.setValue(utils.json_constant.JSONID_FOR_RUNRESULT, utils.json_constant.RUNRESULT_SUCCESS)
			#処理終了。
			return file_path

		except Exception as e:
			utils.config.global_log.error(e)
			jsonObj.setValue(utils.json_constant.JSONID_ERR, "帳票出力に失敗しました")
			jsonObj.setValue(utils.json_constant.JSONID_FOR_RUNRESULT, utils.json_constant.RUNRESULT_FAIL)
			raise
		#utils.config.global_log.debug(str(threading.current_thread().native_id)+ ": end")

	# v_output_f_excelの取得データからExcelファイルを生成する。
	def _generate_excel(self, row) :
		wb = Workbook()
		ws = wb.active
		ws.title = "F相談受付票"

		# タイトル行
		ws.merge_cells("A1:E1")
		title_cell = ws["A1"]
		title_cell.value = "F 相談受付票"
		title_cell.font = Font(size=16, bold=True)
		title_cell.alignment = Alignment(horizontal="center", vertical="center")

		# ヘッダー定義（DB取得項目 → Excel表示ラベル）
		fields = [
			("報告書番号", "報告書番号"),
			("都道府県連", "都道府県連"),
			("商工会", "商工会"),
			("実施日", "実施日"),
			("開始時刻", "開始時刻"),
			("終了時刻", "終了時刻"),
			("支援テーマ", "支援テーマ"),
			("業種", "業種"),
			("事業所名", "事業所名"),
			("担当者名", "担当者名"),
			("担当（主）", "担当（主）"),
			("担当（副）", "担当（副）"),
			("概要", "概要"),
			("内容", "内容"),
		]

		# ヘッダー行（3行目から）
		header_font = Font(bold=True)
		header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
		thin_border = Border(
			left=Side(style="thin"),
			right=Side(style="thin"),
			top=Side(style="thin"),
			bottom=Side(style="thin"),
		)

		for col_idx, (db_key, label) in enumerate(fields, 1):
			cell = ws.cell(row=3, column=col_idx, value=label)
			cell.font = header_font
			cell.fill = header_fill
			cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
			cell.border = thin_border

		# データ行（4行目）
		for col_idx, (db_key, label) in enumerate(fields, 1):
			val = utils.string_util.changeNullToBlank(str(row.get(db_key, "") or ""))
			if db_key == "実施日" and val:
				val = val[:10]
			cell = ws.cell(row=4, column=col_idx, value=val)
			cell.alignment = Alignment(vertical="top", wrap_text=True)
			cell.border = thin_border

		# 列幅調整
		col_widths = [14, 14, 22, 12, 10, 10, 24, 22, 22, 14, 14, 14, 30, 60]
		for i, w in enumerate(col_widths, 1):
			ws.column_dimensions[chr(64 + i)].width = w

		# 内容行の高さを広げる
		ws.row_dimensions[4].height = 120

		# 一時ファイルへ保存する
		temp_dir = tempfile.gettempdir()
		file_name = "F_相談受付票_" + utils.string_util.changeNullToBlank(row.get("報告書番号", "")) + ".xlsx"
		file_path = os.path.join(temp_dir, file_name)
		wb.save(file_path)

		return file_path
