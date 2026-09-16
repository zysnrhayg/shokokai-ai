"""Excel export: v_output_* VIEW + cfg_excel_output_mapping → template."""
import os
import re
import tempfile

import utils.mysqldb_utils as db
import utils.string_util
from app.monthly.services import (
    ANNUAL_AGGREGATE_FORMS,
    MONTHLY_YEARMONTH_FORMS,
    _blank,
    _find_fy,
    _fy_list,
    _g,
    _ym,
    _ensure_dummy_template,
    export_dir,
    is_report_form_code,
    resolve_org,
)

_VIEW_NAME_RE = re.compile(r"^v_output_[a-z0-9_]+$")


def _q(sql, params=None):
    return db.result_to_list_of_dict(db.querySQL(sql, params or {})) or []


def _is_aggregate_form(form_code):
    code = _blank(form_code)
    return code == "H" or code.startswith("I-")


def _as_bool(v):
    if v is True:
        return True
    if v is False or v is None:
        return False
    s = str(v).strip().lower()
    return s in ("1", "t", "true", "yes")


def _get_report_definition(form_code, fiscal_year_id):
    rows = _q(
        """
SELECT form_code, template_filename, view_name, is_day_batch
FROM cfg_excel_report_definition
WHERE deleted_at IS NULL
  AND form_code = :form_code
  AND fiscal_year_id = CAST(:fiscal_year_id AS integer)
LIMIT 1
""",
        {"form_code": form_code, "fiscal_year_id": fiscal_year_id},
    )
    return rows[0] if rows else None


def _get_output_mappings(form_code, fiscal_year_id):
    return _q(
        """
SELECT view_column, sheet_name, cell_address, slot_number
FROM cfg_excel_output_mapping
WHERE deleted_at IS NULL
  AND form_code = :form_code
  AND fiscal_year_id = CAST(:fiscal_year_id AS integer)
ORDER BY slot_number, mapping_id
""",
        {"form_code": form_code, "fiscal_year_id": fiscal_year_id},
    )


def _row_value(row, view_column):
    if not row or not view_column:
        return ""
    if view_column in row:
        return row.get(view_column)
    for key, val in row.items():
        if str(key) == view_column:
            return val
    return ""


def _cell_value(val):
    if val is None:
        return ""
    if hasattr(val, "isoformat"):
        return val.isoformat()
    return utils.string_util.changeNullToBlank(val)


def _fetch_aggregate_row(view_name, prefecture_code, shokokai_cd, fiscal_year_id, year_month=""):
    name = _blank(view_name)
    if not _VIEW_NAME_RE.match(name):
        raise ValueError("不正なVIEW名です")
    params = {
        "prefecture_code": prefecture_code,
        "shokokai_cd": shokokai_cd,
        "fiscal_year_id": fiscal_year_id,
    }
    sql = (
        "SELECT * FROM %s "
        "WHERE prefecture_code = :prefecture_code "
        "AND shokokai_cd = :shokokai_cd "
        "AND fiscal_year_id = CAST(:fiscal_year_id AS integer)"
    ) % name
    ym = _ym(year_month)
    if ym:
        sql += " AND year_month = :year_month"
        params["year_month"] = ym
    sql += " LIMIT 1"
    rows = _q(sql, params)
    return rows[0] if rows else None


def _report_day_for_batch(pref, sho, fy_id, form_code, year_month):
    """日次バッチ: 指定年月内で最終報告日を採用（同一日の複数報告書を1枚に出力）。"""
    rows = _q(
        """
SELECT MAX(r.report_date) AS report_date
FROM trn_report r
WHERE r.prefecture_code = :prefecture_code
  AND r.shokokai_cd = :shokokai_cd
  AND r.fiscal_year_id = CAST(:fiscal_year_id AS integer)
  AND r.form_code = :form_code
  AND r.status = '登録済み'
  AND to_char(r.report_date, 'YYYY-MM') = :year_month
""",
        {
            "prefecture_code": pref,
            "shokokai_cd": sho,
            "fiscal_year_id": fy_id,
            "form_code": form_code,
            "year_month": year_month,
        },
    )
    if not rows or not rows[0].get("report_date"):
        return ""
    rd = rows[0].get("report_date")
    if hasattr(rd, "isoformat"):
        return rd.isoformat()[:10]
    return str(rd)[:10]


def _fetch_report_view_rows(view_name, pref, sho, fy_id, form_code, year_month="", is_day_batch=False, report_id=""):
    name = _blank(view_name)
    if not _VIEW_NAME_RE.match(name):
        raise ValueError("不正なVIEW名です")
    params = {
        "prefecture_code": pref,
        "shokokai_cd": sho,
        "fiscal_year_id": fy_id,
    }
    sql = (
        "SELECT v.* FROM %s v "
        "INNER JOIN trn_report r ON r.report_id = v.report_id "
        "WHERE r.prefecture_code = :prefecture_code "
        "AND r.shokokai_cd = :shokokai_cd "
        "AND r.fiscal_year_id = CAST(:fiscal_year_id AS integer) "
        "AND r.status = '登録済み'"
    ) % name
    if _blank(report_id):
        sql += " AND r.report_id = CAST(:report_id AS integer)"
        params["report_id"] = report_id
    elif is_day_batch:
        ym = _ym(year_month)
        if not ym:
            raise ValueError("実施年月を指定してください")
        report_day = _report_day_for_batch(pref, sho, fy_id, form_code, ym)
        if not report_day:
            return []
        sql += " AND r.report_date = CAST(:report_date AS date)"
        params["report_date"] = report_day
    else:
        ym = _ym(year_month)
        if ym:
            sql += " AND to_char(r.report_date, 'YYYY-MM') = :year_month"
            params["year_month"] = ym
    sql += " ORDER BY r.report_date, r.report_id"
    if is_day_batch:
        sql += " LIMIT 5"
    else:
        sql += " LIMIT 1"
    return _q(sql, params)


def _rows_by_slot(report_rows, is_day_batch):
    """mapping.slot_number ごとに VIEW 行を割り当て。"""
    if not report_rows:
        return {}
    if not is_day_batch:
        row = report_rows[0]
        return {i: row for i in range(0, 6)}
    out = {0: report_rows[0]}
    for idx, row in enumerate(report_rows[:5], start=1):
        out[idx] = row
    return out


def _load_workbook(template_filename):
    from openpyxl import Workbook, load_workbook

    base = export_dir()
    template_name = os.path.basename(_blank(template_filename).replace("\\", "/"))
    if not template_name:
        return Workbook()
    path = os.path.join(base, template_name)
    if os.path.isfile(path):
        return load_workbook(path)
    extracted = _ensure_dummy_template(base, template_name)
    if extracted and os.path.isfile(extracted):
        return load_workbook(extracted)
    return Workbook()


def _requires_yearmonth_for_export(form_code, is_day_batch=False):
    code = _blank(form_code)
    if code in MONTHLY_YEARMONTH_FORMS:
        return True
    if code in ANNUAL_AGGREGATE_FORMS:
        return False
    if is_report_form_code(code):
        return True
    return _as_bool(is_day_batch)


def render_aggregate_excel(session, dto):
    """
    cfg_excel_report_definition.view_name からデータ取得し、
    cfg_excel_output_mapping（slot_number 対応）に従ってテンプレートへ書き込む。
    H/I-* は KPI 集計1行、F/G-* は trn_report（報告書）ベース。
    """
    form_code = _g(dto, "formcode", "form_code")
    if not _blank(form_code):
        raise ValueError("帳票様式が指定されていません")

    _, pref, sho = resolve_org(session, dto)
    if not pref or not sho:
        raise ValueError("組織情報が取得できません")

    fy_list = _fy_list()
    fy_code = _g(dto, "fiscalyearcode", "fiscal_year_code")
    fy = _find_fy(fy_list, code=fy_code) if fy_code else (fy_list[0] if fy_list else None)
    fy_id = (fy or {}).get("fiscal_year_id")
    if fy_id in (None, ""):
        raise ValueError("対象年度が取得できません")
    fy_code_out = _blank((fy or {}).get("fiscal_year_code")) or fy_code

    defn = _get_report_definition(form_code, fy_id)
    if not defn:
        raise FileNotFoundError("帳票の定義がありません（cfg_excel_report_definition）")

    view_name = _blank(defn.get("view_name"))
    template_filename = os.path.basename(_blank(defn.get("template_filename")).replace("\\", "/"))
    is_day_batch = _as_bool(defn.get("is_day_batch"))
    if not view_name or not template_filename:
        raise FileNotFoundError("帳票のVIEWまたはテンプレート定義がありません")

    year_month = _g(dto, "yearmonth", "year_month")
    report_id = _g(dto, "reportid", "report_id")
    if _requires_yearmonth_for_export(form_code, is_day_batch):
        if not _blank(report_id) and not _ym(year_month):
            pass
        elif not _ym(year_month):
            raise ValueError("実施年月を指定してください")

    if _is_aggregate_form(form_code):
        row = _fetch_aggregate_row(
            view_name, pref, sho, fy_id, year_month if form_code in MONTHLY_YEARMONTH_FORMS else ""
        )
        if not row:
            raise FileNotFoundError("出力対象のデータがありません")
        slot_rows = {i: row for i in range(0, 6)}
    else:
        report_rows = _fetch_report_view_rows(
            view_name,
            pref,
            sho,
            fy_id,
            form_code,
            year_month=year_month,
            is_day_batch=is_day_batch,
            report_id=report_id,
        )
        if not report_rows:
            raise FileNotFoundError("出力対象のデータがありません")
        slot_rows = _rows_by_slot(report_rows, is_day_batch)

    mappings = _get_output_mappings(form_code, fy_id)
    if not mappings:
        raise FileNotFoundError("帳票の出力マッピングがありません（cfg_excel_output_mapping）")

    wb = _load_workbook(template_filename)
    for rec in mappings:
        sheet_name = _blank(rec.get("sheet_name")) or "Sheet1"
        cell_addr = _blank(rec.get("cell_address"))
        view_col = _blank(rec.get("view_column"))
        if not cell_addr or not view_col:
            continue
        try:
            slot = int(rec.get("slot_number") if rec.get("slot_number") not in (None, "") else 1)
        except (TypeError, ValueError):
            slot = 1
        row = slot_rows.get(slot)
        if not row:
            continue
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        ws[cell_addr] = _cell_value(_row_value(row, view_col))

    fd, out_path = tempfile.mkstemp(
        suffix=".xlsx",
        prefix="monthly_%s_" % form_code.replace("-", "_"),
    )
    os.close(fd)
    wb.save(out_path)
    return out_path, template_filename, form_code, fy_code_out
