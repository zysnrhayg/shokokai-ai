#import_data_util.py
#py_import_data_util.vm — CSV / Excel import reader (output: utils/import_data_util.py)

import csv
import os
import re

try:
	import openpyxl
except ImportError:
	openpyxl = None

EXPAND_DOWN = "1"
EXPAND_RIGHT = "2"
REQUIRED_MUST = "1"


class ImportDataUtil:

	@staticmethod
	def read_file(file_path, fix_flg, out_put_config, in_out_items=None):
		"""Read import file. fix_flg: 0=CSV, 4=Excel. Extension overrides fix_flg when needed."""
		in_out_items = in_out_items or []
		ext = os.path.splitext(file_path)[1].lower()
		if ext in (".xlsx", ".xlsm"):
			return ImportDataUtil._read_excel(file_path, out_put_config, in_out_items)
		if ext == ".xls":
			raise ValueError("旧形式の.xlsファイルは未対応です。.xlsxまたは.csvをご利用ください。")
		if fix_flg == "0" or ext == ".csv" or not ext:
			ImportDataUtil._ensure_text_csv(file_path)
			return ImportDataUtil.read_csv(
				file_path,
				ImportDataUtil.build_header_map(in_out_items),
				encoding=out_put_config.get("encode") or "UTF-8",
				delimiter=out_put_config.get("divide_char") or ",",
				has_header=str(out_put_config.get("head_flg") or "1") == "1",
			)
		if fix_flg == "4":
			return ImportDataUtil._read_excel(file_path, out_put_config, in_out_items)
		raise ValueError("未対応のインポート形式です。")

	@staticmethod
	def _read_excel(file_path, out_put_config, in_out_items):
		identity_str = (out_put_config.get("identity_str") or "").strip()
		if identity_str:
			field_name_map = ImportDataUtil.build_field_name_map(in_out_items)
			definitions = ImportDataUtil.parse_identity(identity_str, field_name_map)
			return ImportDataUtil.read_excel_by_identity(file_path, definitions)
		return ImportDataUtil.read_excel_by_header(
			file_path,
			ImportDataUtil.build_header_map(in_out_items),
		)

	@staticmethod
	def _display_encoding(encoding):
		enc = (encoding or "UTF-8").strip().lower().replace("_", "-")
		labels = {
			"shift-jis": "日本語(Shift-JIS)",
			"shift_jis": "日本語(Shift-JIS)",
			"sjis": "日本語(Shift-JIS)",
			"cp932": "日本語(CP932)",
			"utf-8": "UTF-8",
			"utf8": "UTF-8",
		}
		return labels.get(enc, encoding or "UTF-8")

	@staticmethod
	def format_read_error(err, expected_encoding=None):
		"""Return a user-facing import read error message (no Python tracebacks)."""
		if isinstance(err, ValueError):
			return str(err)
		if isinstance(err, UnicodeDecodeError):
			enc_label = ImportDataUtil._display_encoding(expected_encoding)
			return (
				"インポートファイルの文字コードが正しくありません。"
				"設定は「" + enc_label + "」です。"
				"Excelの「名前を付けて保存」から「CSV（コンマ区切り）」を選択し、"
				"文字コードを設定に合わせて保存し直してください。"
			)
		if isinstance(err, ImportError):
			return "Excelファイルの読み込みに必要なモジュールが不足しています。システム管理者に連絡ください。"
		return "インポートファイルの読み込みに失敗しました。ファイル形式と文字コードを確認してください。"

	@staticmethod
	def _ensure_text_csv(file_path):
		with open(file_path, "rb") as f:
			head = f.read(8)
		if head.startswith(b"PK"):
			raise ValueError("選択されたファイルはExcel形式です。CSV形式のファイルを指定してください。")
		if b"\x00" in head:
			raise ValueError("インポートファイルの形式が正しくありません。CSV形式のテキストファイルを指定してください。")

	@staticmethod
	def _encoding_decode_error(expected_encoding):
		enc_label = ImportDataUtil._display_encoding(expected_encoding)
		return ValueError(
			"インポートファイルの文字コードが正しくありません。"
			"設定は「" + enc_label + "」です。"
			"Excelの「名前を付けて保存」から「CSV（コンマ区切り）」を選択し、"
			"文字コードを設定に合わせて保存し直してください。"
		)

	@staticmethod
	def _normalize_encoding(encoding):
		enc = (encoding or "UTF-8").strip().lower().replace("_", "-")
		aliases = {
			"shift-jis": "shift_jis",
			"sjis": "shift_jis",
			"cp932": "cp932",
			"utf8": "utf-8",
		}
		return aliases.get(enc, encoding or "UTF-8")

	@staticmethod
	def _encoding_candidates(encoding):
		primary = ImportDataUtil._normalize_encoding(encoding)
		candidates = [primary]
		for fallback in ("cp932", "utf-8-sig", "utf-8"):
			if fallback not in candidates:
				candidates.append(fallback)
		return candidates

	@staticmethod
	def build_field_id_list_from_header_map(header_to_field):
		field_ids = []
		seen = set()
		for key, field_id in header_to_field.items():
			if key == field_id:
				continue
			if field_id not in seen:
				seen.add(field_id)
				field_ids.append(field_id)
		return field_ids

	@staticmethod
	def build_field_label_map(in_out_items):
		labels = {}
		for item in in_out_items or []:
			field_id = item.get("field_id") or item.get("fieldID")
			field_name = item.get("field_name") or item.get("fieldName")
			if field_id and field_name:
				labels[field_id] = field_name
		return labels

	@staticmethod
	def _row_looks_like_header(row, header_to_field):
		if not row:
			return False
		match = 0
		for cell in row:
			if ImportDataUtil._norm_header(cell) in header_to_field:
				match += 1
		return match >= 2

	@staticmethod
	def build_header_map(in_out_items):
		header_map = {}
		for item in in_out_items or []:
			field_id = item.get("field_id") or item.get("fieldID")
			field_name = item.get("field_name") or item.get("fieldName")
			if field_id and field_name:
				header_map[field_name] = field_id
				header_map[field_id] = field_id
		return header_map

	@staticmethod
	def build_field_name_map(in_out_items):
		field_map = {}
		for item in in_out_items or []:
			field_name = item.get("field_name") or item.get("fieldName")
			if field_name:
				field_map[field_name] = item
		return field_map

	@staticmethod
	def read_csv(file_path, header_to_field, encoding="UTF-8", delimiter=",", has_header=True):
		all_rows = None
		last_err = None
		for enc in ImportDataUtil._encoding_candidates(encoding):
			try:
				with open(file_path, "r", encoding=enc, newline="") as f:
					all_rows = list(csv.reader(f, delimiter=delimiter))
				break
			except UnicodeDecodeError as err:
				last_err = err
		if all_rows is None:
			raise ImportDataUtil._encoding_decode_error(encoding)
		if not all_rows:
			return []
		if has_header or ImportDataUtil._row_looks_like_header(all_rows[0], header_to_field):
			return ImportDataUtil._parse_rows(all_rows, header_to_field)
		rows = []
		field_ids = ImportDataUtil.build_field_id_list_from_header_map(header_to_field)
		for data_row in all_rows:
			if not any(str(c or "").strip() for c in data_row):
				continue
			row = {}
			for idx, field_id in enumerate(field_ids):
				row[field_id] = data_row[idx] if idx < len(data_row) else ""
			rows.append(row)
		return rows

	@staticmethod
	def read_excel_by_header(file_path, header_to_field):
		if openpyxl is None:
			raise ImportError("openpyxl is required for Excel import")
		wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
		try:
			ws = wb.active
			all_rows = []
			for row in ws.iter_rows(values_only=True):
				all_rows.append(["" if c is None else str(c).strip() for c in row])
			return ImportDataUtil._parse_rows(all_rows, header_to_field)
		finally:
			wb.close()

	@staticmethod
	def parse_identity(identity_str, field_name_map):
		definitions = []
		for line in identity_str.replace("\r\n", "\n").split("\n"):
			line = line.strip()
			if not line:
				continue
			for part in re.findall(r"\{[^{}]+\}", line):
				definitions.append(ImportDataUtil._parse_identity_entry(part, field_name_map))
		return [d for d in definitions if d.get("field_id")]

	@staticmethod
	def _parse_identity_entry(entry, field_name_map):
		split_re = re.compile(r',(?=(?:[^"]*"[^"]*")*[^"]*$)')
		inner = entry.strip().strip("{}")
		parts = split_re.split(inner)
		clean = [p.strip().strip('"') for p in parts]
		while len(clean) < 7:
			clean.append("")
		item_name = clean[0]
		expand = clean[4]
		if expand in ("下", "DOWN", "down"):
			expand = EXPAND_DOWN
		elif expand in ("右", "RIGHT", "right"):
			expand = EXPAND_RIGHT
		else:
			expand = ""
		required = REQUIRED_MUST if clean[6] in ("必須", "1") else "0"
		item = field_name_map.get(item_name) or {}
		return {
			"item_name": item_name,
			"sheet_name": clean[1],
			"cell_column": int(float(clean[2] or "0")),
			"cell_row": int(float(clean[3] or "0")),
			"expand": expand,
			"property": clean[5],
			"required": required,
			"skip_line_keycode": clean[7] if len(clean) > 7 else "",
			"field_id": item.get("field_id") or item.get("fieldID") or item_name,
		}

	@staticmethod
	def read_excel_by_identity(file_path, definitions):
		if openpyxl is None:
			raise ImportError("openpyxl is required for Excel import")
		wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
		try:
			by_sheet = {}
			for d in definitions:
				by_sheet.setdefault(d["sheet_name"], []).append(d)

			all_rows = []
			for sheet_name, sheet_defs in by_sheet.items():
				ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
				common = [d for d in sheet_defs if not d.get("expand")]
				down = [d for d in sheet_defs if d.get("expand") == EXPAND_DOWN]
				common_values = {}
				for d in common:
					common_values[d["field_id"]] = ImportDataUtil._get_cell_value(
						ws, d["cell_row"], d["cell_column"]
					)
				if down:
					min_row = min(d["cell_row"] for d in down)
					max_row = ws.max_row or min_row
					for row_idx in range(min_row, max_row + 1):
						row = dict(common_values)
						skip = False
						for d in down:
							val = ImportDataUtil._get_cell_value(ws, row_idx, d["cell_column"])
							if ImportDataUtil._should_skip_line(val, d.get("skip_line_keycode")):
								skip = True
								break
							row[d["field_id"]] = val
						if skip:
							continue
						if any(str(v or "").strip() for v in row.values()):
							all_rows.append(row)
				elif common_values:
					all_rows.append(common_values)
			return all_rows
		finally:
			wb.close()

	@staticmethod
	def _should_skip_line(value, skip_keys):
		if not skip_keys:
			return False
		text = str(value or "")
		for key in skip_keys.split(","):
			key = key.strip()
			if key and key in text:
				return True
		return not str(value or "").strip() and "空白" in skip_keys

	@staticmethod
	def _get_cell_value(ws, row_index, col_index):
		# design book uses 1-based row/col
		cell = ws.cell(row=row_index, column=col_index)
		val = cell.value
		if val is None:
			return ""
		if hasattr(val, "strftime"):
			return val.strftime("%Y/%m/%d")
		return str(val).strip()

	@staticmethod
	def _parse_rows(all_rows, header_to_field):
		if len(all_rows) < 2:
			return []
		header = [ImportDataUtil._norm_header(c) for c in all_rows[0]]
		col_map = {}
		for idx, label in enumerate(header):
			field = header_to_field.get(label)
			if field:
				col_map[field] = idx
		rows = []
		for data_row in all_rows[1:]:
			if not any(str(c or "").strip() for c in data_row):
				continue
			row = {}
			for field, idx in col_map.items():
				row[field] = data_row[idx] if idx < len(data_row) else ""
			rows.append(row)
		return rows

	@staticmethod
	def _norm_header(cell):
		if cell is None:
			return ""
		return str(cell).strip()

	@staticmethod
	def validate_row(row_num, row, required_fields, property_map=None, field_labels=None):
		property_map = property_map or {}
		field_labels = field_labels or {}
		for field in required_fields:
			if not str(row.get(field, "") or "").strip():
				label = field_labels.get(field, field)
				return f"{row_num}行目の「{label}」列は必須項目のため、空白で設定はできません。"
		for field, prop in property_map.items():
			val = str(row.get(field, "") or "").strip()
			if not val:
				continue
			label = field_labels.get(field, field)
			if prop in ("数字", "numeric", "number") and not re.match(r"^(0|\+?[1-9][0-9]*)$", val.replace(",", "")):
				return f"{row_num}行目の「{label}」列には数字のみ入力可能です。"
			if prop in ("日付", "date") and not re.match(r"^\d{4}[-/]\d{2}[-/]\d{2}$", val):
				return f"{row_num}行目の「{label}」列には日付のみ入力可能です。"
			if prop in ("日時", "datetime") and not re.match(r"^\d{4}[-/]\d{2}[-/]\d{2}([ T]\d{2}:\d{2}:\d{2})?$", val):
				return f"{row_num}行目の「{label}」列には日時のみ入力可能です。[yyyy/mm/dd hh:mm:ss]の形式で入力してください。"
		return ""

	@staticmethod
	def resolve_import_datetime(value, default_now):
		"""Normalize CSV datetime: blank -> default_now; date-only -> 00:00:00; invalid -> None."""
		text = str(value or "").strip()
		if not text:
			return default_now
		import utils.date_util
		dt = utils.date_util.parse_flexible_datetime(text)
		if dt is None:
			return None
		return dt.strftime("%Y-%m-%d %H:%M:%S")
